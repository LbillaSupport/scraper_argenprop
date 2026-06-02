"""
exporter.py
===========
Genera el archivo Excel final con varias hojas de ranking.

Hojas:
  TODAS          -> todas las propiedades terminadas (sin pozo)
  OPORTUNIDADES  -> USD/m² cubierto ascendente
  MAYOR_M2       -> m² cubiertos descendente
  MENOR_PRECIO   -> precio ascendente
  BAJAS_EXPENSAS -> expensas ascendente
  SCORE          -> score descendente
  EN_POZO        -> todas las detectadas en pozo

Formato:
  * Autofiltro en cada hoja.
  * Primera fila congelada.
  * Ancho de columnas ajustado al contenido.
  * Formato moneda para Precio y Expensas.
  * Formato numérico para m² y USD/m².
"""

from __future__ import annotations

import os
from datetime import datetime

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# Orden y etiquetas de las columnas del Excel (clave interna -> encabezado)
COLUMNS: list[tuple[str, str]] = [
    ("url", "URL"),
    ("neighborhood", "Barrio"),
    ("address", "Dirección"),
    ("title", "Título"),
    ("property_kind", "Tipo"),
    ("price", "Precio"),
    ("currency", "Moneda"),
    ("expenses", "Expensas"),
    ("expenses_est", "Expensas est."),
    ("m2_covered", "M² Cubiertos"),
    ("m2_total", "M² Totales"),
    ("rooms", "Ambientes"),
    ("bedrooms", "Dormitorios"),
    ("bathrooms", "Baños"),
    ("age", "Antigüedad"),
    ("usd_m2_covered", "$/m² Cub"),
    ("usd_m2_total", "$/m² Total"),
    ("zone_score", "Punt_Zona"),
    ("score", "Score"),
    ("en_pozo", "En_Pozo"),
    ("description", "Descripción"),
]

HEADERS = [h for _, h in COLUMNS]

# Formatos numéricos por encabezado
CURRENCY_COLS = {"Precio", "Expensas", "Expensas est."}
NUMBER_COLS = {
    "M² Cubiertos",
    "M² Totales",
    "$/m² Cub",
    "$/m² Total",
    "Punt_Zona",
    "Score",
}
CURRENCY_FMT = '"$"#,##0'
NUMBER_FMT = "#,##0.00"

# Etiquetas legibles para el tipo de propiedad detectado.
KIND_LABELS = {"departamento": "Depto", "ph": "PH", "casa": "Casa"}

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
HYPERLINK_FONT = Font(color="0563C1", underline="single")


def _cell_value(key: str, r: dict):
    """Formatea el valor de una celda según la columna."""
    if key == "en_pozo":
        return "Sí" if r.get(key) else "No"
    if key == "property_kind":
        return KIND_LABELS.get(r.get(key), r.get(key))
    return r.get(key)


def _records_to_df(records: list[dict]) -> pd.DataFrame:
    """Convierte los registros a un DataFrame con las columnas finales."""
    rows = [{header: _cell_value(key, r) for key, header in COLUMNS} for r in records]
    df = pd.DataFrame(rows, columns=HEADERS)
    return df


def _sorted(records: list[dict], key: str, ascending: bool) -> list[dict]:
    """Ordena dejando los None al final."""
    inf = float("inf")
    def sort_key(r):
        v = r.get(key)
        if v is None:
            return inf if ascending else -inf
        return v
    return sorted(records, key=sort_key, reverse=not ascending)


def export(ok: list[dict], pozo: list[dict], path: str | None = None) -> str:
    """
    Genera el Excel y devuelve la ruta del archivo creado.
    """
    if path is None:
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        path = f"argenprop_{stamp}.xlsx"

    # Aseguramos que exista la carpeta destino.
    parent = os.path.dirname(os.path.abspath(path))
    os.makedirs(parent, exist_ok=True)

    # Definición de cada hoja: (nombre, registros ya ordenados)
    sheets: dict[str, list[dict]] = {
        "TODAS": ok,
        "OPORTUNIDADES": _sorted(ok, "usd_m2_covered", ascending=True),
        "MAYOR_M2": _sorted(ok, "m2_covered", ascending=False),
        "MENOR_PRECIO": _sorted(ok, "price", ascending=True),
        "BAJAS_EXPENSAS": _sorted(ok, "expenses", ascending=True),
        "SCORE": _sorted(ok, "score", ascending=False),
        "EN_POZO": pozo,
    }

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for sheet_name, recs in sheets.items():
            df = _records_to_df(recs)
            # Si no hay datos igual escribimos los encabezados.
            if df.empty:
                df = pd.DataFrame(columns=HEADERS)
            df.to_excel(writer, sheet_name=sheet_name, index=False)
            _format_sheet(writer.book[sheet_name], df)

    return path


def _format_sheet(ws, df: pd.DataFrame) -> None:
    """Aplica formato a una hoja: encabezados, filtros, congelado, anchos."""
    n_rows, n_cols = df.shape

    # --- Encabezados con estilo ---
    for col_idx, header in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # --- URLs como hipervínculos reales (un solo click, sin editar la celda) ---
    if "URL" in HEADERS:
        url_col = HEADERS.index("URL") + 1
        for row in range(2, n_rows + 2):
            cell = ws.cell(row=row, column=url_col)
            value = cell.value
            if isinstance(value, str) and value.startswith("http"):
                cell.hyperlink = value
                cell.font = HYPERLINK_FONT

    # --- Congelar primera fila ---
    ws.freeze_panes = "A2"

    # --- Autofiltro (incluye encabezado + datos) ---
    last_col_letter = get_column_letter(n_cols)
    ws.auto_filter.ref = f"A1:{last_col_letter}{max(n_rows + 1, 1)}"

    # --- Formatos numéricos por columna ---
    for col_idx, header in enumerate(HEADERS, start=1):
        if header in CURRENCY_COLS:
            fmt = CURRENCY_FMT
        elif header in NUMBER_COLS:
            fmt = NUMBER_FMT
        else:
            fmt = None
        if fmt:
            for row in range(2, n_rows + 2):
                ws.cell(row=row, column=col_idx).number_format = fmt

    # --- Ancho de columnas ajustado al contenido (con tope) ---
    for col_idx, header in enumerate(HEADERS, start=1):
        max_len = len(str(header))
        for row in range(2, n_rows + 2):
            value = ws.cell(row=row, column=col_idx).value
            if value is not None:
                max_len = max(max_len, len(str(value)))
        # La descripción y la URL pueden ser enormes: las acotamos.
        cap = 60 if header in {"Descripción", "URL", "Título", "Dirección"} else 22
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, cap)
